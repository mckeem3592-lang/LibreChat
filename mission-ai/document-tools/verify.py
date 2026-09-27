from importlib.metadata import version
from pathlib import Path
from tempfile import TemporaryDirectory

EXPECTED = {
    "openpyxl": "3.1.5",
    "python-docx": "1.2.0",
    "python-pptx": "1.0.2",
    "pypdf": "6.19.0",
    "reportlab": "5.0.1",
}

for package, expected in EXPECTED.items():
    actual = version(package)
    if actual != expected:
        raise SystemExit(f"{package}: expected {expected}, got {actual}")

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas


def verify_xlsx(root: Path) -> None:
    path = root / "verify.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Mission AI"
    sheet["B1"] = 42
    workbook.save(path)

    reopened = load_workbook(path, data_only=False)
    assert reopened.active["A1"].value == "Mission AI"
    assert reopened.active["B1"].value == 42
    reopened.close()


def verify_docx(root: Path) -> None:
    path = root / "verify.docx"
    document = Document()
    document.add_heading("Mission AI", level=1)
    document.add_paragraph("Document round-trip")
    document.save(path)

    reopened = Document(path)
    text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    assert "Mission AI" in text
    assert "Document round-trip" in text


def verify_pptx(root: Path) -> None:
    path = root / "verify.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Mission AI"
    slide.placeholders[1].text = "Presentation round-trip"
    presentation.save(path)

    reopened = Presentation(path)
    assert len(reopened.slides) == 1
    texts = [
        shape.text
        for shape in reopened.slides[0].shapes
        if hasattr(shape, "text")
    ]
    assert any("Mission AI" in value for value in texts)


def verify_pdf(root: Path) -> None:
    source = root / "verify.pdf"
    copied = root / "verify-copy.pdf"

    pdf = canvas.Canvas(str(source))
    pdf.drawString(72, 720, "Mission AI PDF round-trip")
    pdf.showPage()
    pdf.save()

    reader = PdfReader(str(source))
    assert len(reader.pages) == 1

    writer = PdfWriter()
    writer.add_page(reader.pages[0])
    with copied.open("wb") as handle:
        writer.write(handle)

    reopened = PdfReader(str(copied))
    assert len(reopened.pages) == 1


with TemporaryDirectory(prefix="mission-ai-doc-verify-") as temp:
    root = Path(temp)
    verify_xlsx(root)
    verify_docx(root)
    verify_pptx(root)
    verify_pdf(root)

print("Mission AI document toolchain verified with XLSX/DOCX/PPTX/PDF round trips.")
