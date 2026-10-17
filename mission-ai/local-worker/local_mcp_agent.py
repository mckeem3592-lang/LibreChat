"""Owner-scoped MCP worker for local documents and allowlisted browser actions.

Queued work survives an MCP request disconnect, but an OS process restart interrupts it.
Status polling uses new requests; it does not extend the original HTTP request.
Restart the worker to load source changes; in-flight tasks do not survive that restart.
"""

from __future__ import annotations

import asyncio
from concurrent.futures import ThreadPoolExecutor
import gc
import hmac
import ipaddress
import json
import math
import os
import re
import secrets
import socket
import stat
import tempfile
import time
from pathlib import Path
from typing import Any, Awaitable, Callable
from urllib.parse import urlsplit

from docx import Document
from mcp.server import MCPServer
from mcp.server.transport_security import TransportSecuritySettings
from openpyxl import Workbook, load_workbook
from playwright.async_api import async_playwright
import pymupdf
from starlette.responses import PlainTextResponse
import uvicorn


TOKEN = os.environ.get("LOCAL_MCP_TOKEN", "")
OWNER_ID = os.environ.get("LOCAL_MCP_OWNER_USER_ID", "")
ROOT_SETTING = os.environ.get("LOCAL_MCP_ROOT", "")
PUBLIC_URL = os.environ.get("LOCAL_MCP_PUBLIC_URL", "")
ORIGINS_SETTING = os.environ.get("LOCAL_MCP_ALLOWED_ORIGINS", "")

if (len(TOKEN) < 40 or not TOKEN.isascii() or any(char.isspace() for char in TOKEN)
        or not re.fullmatch(r"[0-9a-fA-F]{24}", OWNER_ID)
        or not ROOT_SETTING or not PUBLIC_URL or not ORIGINS_SETTING):
    raise RuntimeError("LOCAL_MCP_TOKEN, LOCAL_MCP_OWNER_USER_ID, LOCAL_MCP_ROOT, LOCAL_MCP_PUBLIC_URL and LOCAL_MCP_ALLOWED_ORIGINS are required")

ROOT = Path(ROOT_SETTING).expanduser().resolve(strict=True)
if not ROOT.is_dir():
    raise RuntimeError("LOCAL_MCP_ROOT must be an existing directory")

public = urlsplit(PUBLIC_URL)
if public.scheme != "https" or not public.hostname or public.username or public.password or public.path.rstrip("/") != "/mcp" or public.query or public.fragment:
    raise RuntimeError("LOCAL_MCP_PUBLIC_URL must be an HTTPS URL ending in /mcp")

ALLOWED_ORIGINS = frozenset(value.strip().rstrip("/") for value in ORIGINS_SETTING.split(",") if value.strip())
if not ALLOWED_ORIGINS or any(urlsplit(value).scheme != "https" or not urlsplit(value).hostname or
                              urlsplit(value).port not in (None, 443) or
                              urlsplit(value).username or urlsplit(value).password or
                              urlsplit(value).path or urlsplit(value).query or urlsplit(value).fragment
                              for value in ALLOWED_ORIGINS):
    raise RuntimeError("LOCAL_MCP_ALLOWED_ORIGINS must contain HTTPS origins only")


def scoped_path(value: str, *, must_exist: bool) -> Path:
    if not value or "\x00" in value:
        raise ValueError("Invalid path")
    given = Path(value).expanduser()
    candidate = given if given.is_absolute() else ROOT / given
    resolved = candidate.resolve(strict=must_exist)
    if resolved != ROOT and ROOT not in resolved.parents:
        raise ValueError("Path is outside LOCAL_MCP_ROOT")
    probe = candidate
    while probe != ROOT and ROOT in probe.parents:
        if probe.is_symlink():
            raise ValueError("Symbolic links are not allowed")
        probe = probe.parent
    return resolved


def existing_file(value: str, suffix: str | None = None) -> Path:
    path = scoped_path(value, must_exist=True)
    if not path.is_file() or (suffix and path.suffix.lower() != suffix):
        raise ValueError("Expected a supported regular file")
    if path.stat().st_size > 20_000_000:
        raise ValueError("File exceeds 20 MB")
    return path


def atomic_save(save: Any, target: Path) -> None:
    if not target.parent.is_dir():
        raise ValueError("Destination parent must already exist within LOCAL_MCP_ROOT")
    fd, temporary = tempfile.mkstemp(prefix=".mission-ai-", suffix=target.suffix, dir=target.parent)
    os.close(fd)
    try:
        if target.exists():
            os.chmod(temporary, stat.S_IMODE(target.stat().st_mode))
        save(temporary)
        if Path(temporary).stat().st_size > 20_000_000:
            raise ValueError("Result exceeds 20 MB")
        os.replace(temporary, target)
    finally:
        Path(temporary).unlink(missing_ok=True)


def relative_result(path: Path) -> str:
    return str(path.relative_to(ROOT))


def public_site(url: str) -> bool:
    parsed = urlsplit(url)
    origin = f"{parsed.scheme}://{parsed.netloc}".rstrip("/")
    if (parsed.scheme != "https" or origin not in ALLOWED_ORIGINS or not parsed.hostname
            or parsed.port not in (None, 443) or parsed.username or parsed.password):
        return False
    try:
        addresses = socket.getaddrinfo(parsed.hostname, parsed.port or 443,
                                        family=socket.AF_INET, type=socket.SOCK_STREAM)
        return bool(addresses) and all(ipaddress.ip_address(result[4][0]).is_global for result in addresses)
    except (OSError, ValueError):
        return False


mcp = MCPServer("mission-ai-local-documents")

AUTOMATION_LOG = Path.home() / "Library" / "Logs" / "MissionAI" / "local-mcp-automation.log"
AUTOMATION_RESULTS: dict[str, dict[str, Any]] = {}
AUTOMATION_TASKS: dict[str, asyncio.Task[None]] = {}
AUTOMATION_LIMIT = asyncio.Semaphore(1)  # Serialize edits so two jobs cannot overwrite the same file.
AUTOMATION_THREADS = ThreadPoolExecutor(max_workers=4, thread_name_prefix="mission-ai-automation")
MAX_PENDING_AUTOMATIONS = 16
MAX_AUTOMATION_RESULTS = 128
RESULT_TTL_SECONDS = 3600


def log_automation(task_id: str, tool: str, status: str, error: str | None = None) -> None:
    """Record lifecycle metadata only; never write file contents or browser text."""
    AUTOMATION_LOG.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
    flags = os.O_WRONLY | os.O_APPEND | os.O_CREAT | getattr(os, "O_NOFOLLOW", 0)
    descriptor = os.open(AUTOMATION_LOG, flags, 0o600)
    try:
        if not stat.S_ISREG(os.fstat(descriptor).st_mode):
            raise OSError("Automation log is not a regular file")
        os.fchmod(descriptor, 0o600)
        event = {"task_id": task_id, "tool": tool, "status": status, "at": time.time()}
        if error:
            event["error"] = error
        os.write(descriptor, (json.dumps(event, separators=(",", ":")) + "\n").encode("utf-8"))
    finally:
        os.close(descriptor)


def run_file_operation(operation: Callable[[], Any]) -> Any:
    try:
        return operation()
    finally:
        # The document libraries can retain reference cycles after a completed edit.
        gc.collect()


async def run_automation(
    task_id: str, tool: str, operation: Callable[[], Awaitable[Any]],
) -> None:
    state = AUTOMATION_RESULTS[task_id]
    try:
        await asyncio.to_thread(log_automation, task_id, tool, "queued")
        async with AUTOMATION_LIMIT:
            state.update(status="processing", step=2)
            await asyncio.to_thread(log_automation, task_id, tool, "processing")
            result = await operation()
            state.update(status="completed", step=3, result=result)
    except asyncio.CancelledError:
        state.update(status="failed", step=3, error="worker_stopped")
        raise
    except Exception as error:
        state.update(status="failed", step=3, error=type(error).__name__,
                     message=str(error)[:300])
    finally:
        state["finished_at"] = time.monotonic()
        try:
            await asyncio.to_thread(
                log_automation, task_id, tool, state["status"], state.get("error"),
            )
        except OSError:
            pass  # The result remains queryable even if the audit log becomes unavailable.
        AUTOMATION_TASKS.pop(task_id, None)


def start_automation(tool: str, operation: Callable[[], Awaitable[Any]]) -> dict[str, Any]:
    """Return a receipt immediately; the worker event loop owns the operation."""
    asyncio.get_running_loop().set_default_executor(AUTOMATION_THREADS)
    now = time.monotonic()
    for task_id, state in list(AUTOMATION_RESULTS.items()):
        if now - state.get("finished_at", now) > RESULT_TTL_SECONDS:
            AUTOMATION_RESULTS.pop(task_id, None)
    while len(AUTOMATION_RESULTS) >= MAX_AUTOMATION_RESULTS:
        oldest_finished = next((key for key in AUTOMATION_RESULTS
                                if key not in AUTOMATION_TASKS), None)
        if oldest_finished is None:
            break
        AUTOMATION_RESULTS.pop(oldest_finished)
    if len(AUTOMATION_TASKS) >= MAX_PENDING_AUTOMATIONS:
        raise RuntimeError("Automation queue is full; retry later")
    task_id = secrets.token_hex(16)
    AUTOMATION_RESULTS[task_id] = {"status": "processing", "step": 1}
    AUTOMATION_TASKS[task_id] = asyncio.create_task(run_automation(task_id, tool, operation))
    return {"status": "processing", "task_id": task_id, "step": 1, "poll_after_seconds": 15}


@mcp.tool()
async def check_automation_status(task_id: str) -> dict[str, Any]:
    """Poll a submitted local automation after about 15 seconds for its result."""
    if not re.fullmatch(r"[0-9a-f]{32}", task_id):
        raise ValueError("Invalid task ID")
    state = AUTOMATION_RESULTS.get(task_id)
    if state is None:
        return {"task_id": task_id, "status": "unknown",
                "message": "Task not found, expired, or worker restarted"}
    return {"task_id": task_id, **{key: value for key, value in state.items()
                                   if key in {"status", "step", "result", "error", "message"}}}


@mcp.tool()
async def modify_word_doc(path: str, text_to_append: str) -> dict[str, Any]:
    """Queue a Word edit; poll check_automation_status with the returned task_id."""
    if not text_to_append or len(text_to_append) > 100_000:
        raise ValueError("Text must contain 1 to 100000 characters")

    def edit() -> str:
        target = existing_file(path, ".docx")
        document = Document(target)
        document.add_paragraph(text_to_append)
        atomic_save(document.save, target)
        return relative_result(target)

    return start_automation("modify_word_doc", lambda: asyncio.to_thread(run_file_operation, edit))


@mcp.tool()
async def write_excel_sheet(path: str, matrix_data: list[list[str | int | float | bool | None]]) -> dict[str, Any]:
    """Queue an Excel edit; poll check_automation_status with the returned task_id."""
    if not isinstance(matrix_data, list) or len(matrix_data) > 1000 or any(
        not isinstance(row, list) or len(row) > 100 for row in matrix_data
    ) or sum(map(len, matrix_data)) > 20_000:
        raise ValueError("Matrix exceeds 1000 rows, 100 columns, or 20000 cells")
    for row in matrix_data:
        for value in row:
            if value is not None and not isinstance(value, (str, int, float, bool)):
                raise ValueError("Matrix cells must be JSON scalars")
            if isinstance(value, str) and len(value) > 32_767:
                raise ValueError("Cell text is too long")
            if isinstance(value, float) and not math.isfinite(value):
                raise ValueError("Cell numbers must be finite")

    def edit() -> str:
        target = scoped_path(path, must_exist=False)
        if target.suffix.lower() != ".xlsx":
            raise ValueError("Only .xlsx workbooks are supported")
        if target.exists():
            existing_file(path, ".xlsx")
            workbook = load_workbook(target)
        else:
            workbook = Workbook()
        sheet = workbook.active
        for row_index, row in enumerate(matrix_data, 1):
            for column_index, value in enumerate(row, 1):
                cell = sheet.cell(row_index, column_index)
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = "s"
        atomic_save(workbook.save, target)
        return relative_result(target)

    return start_automation("write_excel_sheet", lambda: asyncio.to_thread(run_file_operation, edit))


@mcp.tool()
async def modify_pdf_text(path: str, search_text: str, replace_text: str) -> dict[str, Any]:
    """Queue a PDF text replacement; poll check_automation_status for the result."""
    if (not search_text or len(search_text) > 500 or not replace_text
            or len(replace_text) > 500 or "\n" in replace_text or "\r" in replace_text):
        raise ValueError("Search and replacement must be single-line text of 1 to 500 characters")
    if any(ord(char) < 32 or ord(char) > 126 for char in replace_text):
        raise ValueError("PDF replacement currently supports printable ASCII text only")

    def edit() -> dict[str, str | int]:
        target = existing_file(path, ".pdf")
        document = pymupdf.open(target)
        try:
            if document.is_encrypted or document.page_count > 200:
                raise ValueError("Encrypted PDFs or PDFs over 200 pages are not supported")
            matches = []
            for page in document:
                for rect in page.search_for(search_text):
                    matches.append((page.number, rect))
                    if len(matches) > 100:
                        raise ValueError("More than 100 PDF matches are not supported")
            if not matches:
                raise ValueError("Search text was not found as selectable PDF text")
            for page_index, rect in matches:
                page = document[page_index]
                page.add_redact_annot(rect, fill=(1, 1, 1), cross_out=False)
            for page_index in set(index for index, _ in matches):
                document[page_index].apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                                                      graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
            for page_index, rect in matches:
                page = document[page_index]
                box = pymupdf.Rect(rect.x0, rect.y0 - 2, max(rect.x1, rect.x0 + 10), rect.y1 + 3)
                if page.insert_textbox(box, replace_text, fontname="helv", fontsize=8,
                                       color=(0, 0, 0)) < 0:
                    raise ValueError("Replacement text does not fit the original PDF area")

            def save(temporary: str) -> None:
                document.save(temporary, garbage=4, deflate=True)
                document.close()

            atomic_save(save, target)
            return {"path": relative_result(target), "replacements": len(matches)}
        finally:
            if not document.is_closed:
                document.close()

    return start_automation("modify_pdf_text", lambda: asyncio.to_thread(run_file_operation, edit))


@mcp.tool()
async def move_local_file(source: str, destination: str) -> dict[str, Any]:
    """Queue a scoped file move; poll check_automation_status for the result."""

    def move() -> str:
        current = existing_file(source)
        target = scoped_path(destination, must_exist=False)
        if target.exists() or target.is_symlink():
            raise FileExistsError("Destination already exists")
        if not target.parent.is_dir():
            raise ValueError("Destination parent must already exist within LOCAL_MCP_ROOT")
        os.link(current, target, follow_symlinks=False)
        try:
            current.unlink()
        except Exception:
            target.unlink(missing_ok=True)
            raise
        return relative_result(target)

    return start_automation("move_local_file", lambda: asyncio.to_thread(run_file_operation, move))


async def run_local_browser_control(url: str, action: str, selector: str = "", text: str = "") -> dict[str, str]:
    """Navigate, click, fill, or extract text on an allowlisted public HTTPS site."""
    if (len(url) > 2048 or action not in {"navigate", "click", "fill", "extract"}
            or len(selector) > 500 or len(text) > 10_000
            or (action in {"click", "fill"} and not selector)
            or (action == "fill" and not text)
            or not await asyncio.to_thread(public_site, url)):
        raise ValueError("URL is not an allowed public HTTPS site")
    pinned = {}
    for origin in ALLOWED_ORIGINS:
        host = urlsplit(origin).hostname
        addresses = await asyncio.to_thread(socket.getaddrinfo, host, 443, socket.AF_INET,
                                            socket.SOCK_STREAM)
        ips = {item[4][0] for item in addresses}
        if not ips or not all(ipaddress.ip_address(ip).is_global for ip in ips):
            raise ValueError("Allowed origin did not resolve to public IPv4 addresses")
        pinned[host] = sorted(ips)[0]
    resolver_rules = ", ".join(f"MAP {host} {address}" for host, address in pinned.items())
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(headless=True,
                                                   args=[f"--host-resolver-rules={resolver_rules}"])
        try:
            context = await browser.new_context(accept_downloads=False)
            page = await context.new_page()

            async def route_request(route: Any) -> None:
                permitted = await asyncio.to_thread(public_site, route.request.url)
                if permitted:
                    await route.continue_()
                else:
                    await route.abort()

            await context.route("**/*", route_request)
            await page.goto(url, wait_until="domcontentloaded", timeout=15_000)
            if action in {"click", "fill", "extract"}:
                locator = page.locator(selector or "body")
                if action == "click":
                    await locator.click(timeout=10_000)
                elif action == "fill":
                    await locator.fill(text, timeout=10_000)
            final_url = page.url
            if not await asyncio.to_thread(public_site, final_url):
                raise ValueError("Navigation left the allowed origin")
            title = (await page.title())[:300]
            extracted = (await page.locator(selector or "body").inner_text(timeout=10_000))[:20_000] \
                if action == "extract" else (await page.locator("body").inner_text(timeout=10_000))[:20_000]
            return {"url": final_url, "title": title, "text": extracted, "action": action}
        finally:
            await browser.close()


@mcp.tool()
async def local_browser_control(
    url: str, action: str, selector: str = "", text: str = "",
) -> dict[str, Any]:
    """Queue an allowlisted browser action; poll check_automation_status for the result."""
    if (len(url) > 2048 or action not in {"navigate", "click", "fill", "extract"}
            or len(selector) > 500 or len(text) > 10_000
            or (action in {"click", "fill"} and not selector)
            or (action == "fill" and not text)):
        raise ValueError("Invalid browser action")
    return start_automation(
        "local_browser_control",
        lambda: run_local_browser_control(url, action, selector, text),
    )


class BearerAuth:
    def __init__(self, application: Any) -> None:
        self.application = application

    async def __call__(self, scope: dict[str, Any], receive: Any, send: Any) -> None:
        if scope["type"] == "http":
            values = [value for key, value in scope.get("headers", []) if key.lower() == b"authorization"]
            owners = [value for key, value in scope.get("headers", []) if key.lower() == b"x-librechat-user-id"]
            expected = f"Bearer {TOKEN}".encode("ascii")
            if len(values) != 1 or len(owners) != 1 or not hmac.compare_digest(values[0], expected) or not hmac.compare_digest(owners[0], OWNER_ID.encode("ascii")):
                await PlainTextResponse("Unauthorized", status_code=401)(scope, receive, send)
                return
        await self.application(scope, receive, send)


security = TransportSecuritySettings(
    allowed_hosts=[public.netloc, f"{public.hostname}:443", "127.0.0.1:5000", "localhost:5000"],
    allowed_origins=[],
)
app = BearerAuth(mcp.streamable_http_app(transport_security=security))

if __name__ == "__main__":
    uvicorn.run(app, host="127.0.0.1", port=5000, proxy_headers=False, access_log=False)
